"""Update poster/a0/poster_A0_v1.pptx to the aux-pretrained, weight-averaged MTM (test AUC 0.738).

The model is `twoe` in scripts/boost_mtm.py: 3 seeds, each scored by the mean logit of its 5 best
development-AUC epochs. Scores come from the per-epoch logits saved in artifacts/pw_1.5_mtm_boost/runs.

python poster/update_v1.py score     per-user records and panel stats -> poster/panel_stats_v1.json
python poster/update_v1.py draw      panels -> poster/a0/v1/*.svg
python poster/update_v1.py patch     swap the panels and texts inside poster_A0_v1.pptx
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
POSTER = Path(__file__).resolve().parent
OUT = POSTER / "a0" / "v1"
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(POSTER))

import boost_mtm as boost  # noqa: E402
from make_panels import AUX_TARGETS, best_f1_threshold, pick_examples, sigmoid  # noqa: E402

VARIANT = "twoe"
# Keep the original quotes when that user is still in the same cell. The old false-alarm
# user is now scored correctly, so that cell uses a new user and a new excerpt.
V1_TEXTS = {
    "False alarm": [
        "I'm still in quite the depressive funk, and only haven't left work early because I left early yesterday.",
        "Guess I wasn't kidding when I said \"it's all over but the crying\". That's all I want to do right now. "
        "I feel broken. Defective.",
    ],
}


def fold_users():
    from ssr.config import load_config
    from ssr.data.cohort import build_cohort
    from ssr.train.cv import _make_splits

    cfg = load_config(ROOT / "eli_matrix.yaml")
    cohort = build_cohort(cfg, assert_paper=False)
    splits = _make_splits(cohort["y_high"].to_numpy(), 5, cfg.seed, float(cfg.train["train_frac"]),
                          float(cfg.train["dev_frac"]))
    return cohort, [(fold_i, te) for fold_i, _tr, _dv, te in splits]


def ensemble(runs: dict, fold_i: int):
    """Mean raw logit over seeds of each seed's 5 best development-AUC epochs, and the matching heads."""
    dv = np.mean([boost.per_fold_scores(runs[s][fold_i], "snap")[0] for s in runs], 0)
    te = np.mean([boost.per_fold_scores(runs[s][fold_i], "snap")[1] for s in runs], 0)
    heads = np.mean([boost.heads_of(runs[s][fold_i])[1] for s in runs], 0)
    return dv, te, heads


def records(cohort, folds, which: str):
    from ssr.train.metrics import auc_roc

    runs = boost.load_runs(VARIANT)
    rows, aux_fold, aucs = [], [], []
    for fold_i, te in folds:
        c = np.load(boost.CACHE / f"fold{fold_i}.npz")
        frame = cohort.iloc[te]
        assert np.array_equal(frame["y_high"].to_numpy(), c["yte"].astype(int)), "fold order differs from the cache"
        if which == "old":
            l_dv, l_te = c["mtm_dv"] + 0.1 * c["phq_dv"], c["mtm_te"] + 0.1 * c["phq_te"]
        else:
            l_dv, l_te, heads = ensemble(runs, fold_i)
            r = boost.corr_cols(heads, c["zte"])
            aux_fold.append({"fold": fold_i, **{col: float(r[boost.HEADS.index(col)]) for col in AUX_TARGETS}})
        p_dv, p_te = sigmoid(l_dv), sigmoid(l_te)
        thr = best_f1_threshold(c["ydv"].astype(int), p_dv)
        y = c["yte"].astype(int)
        aucs.append(float(auc_roc(y, p_te)))
        for i in range(len(y)):
            rows.append({"fold": fold_i, "uid": frame.iloc[i]["UserId"], "suicide": int(frame.iloc[i]["suicide"]),
                         "y": int(y[i]), "prob": float(p_te[i]), "pred": int(p_te[i] >= thr), "threshold": thr})
    rec = pd.DataFrame(rows)
    rec["correct"] = (rec["pred"] == rec["y"]).astype(int)
    return rec, aux_fold, aucs


def score() -> None:
    cohort, folds = fold_users()
    old, _, old_aucs = records(cohort, folds, "old")
    new, aux_fold, aucs = records(cohort, folds, "new")
    print(f"old MTM + PHQ-9: AUC {np.mean(old_aucs):.4f}  accuracy {old.correct.mean():.3f}")
    print(f"new MTM:         AUC {np.mean(aucs):.4f}  {[round(a, 4) for a in aucs]}  accuracy {new.correct.mean():.3f}")
    old_pick, new_pick = pick_examples(old), pick_examples(new)
    by_uid = new.set_index("uid")
    chosen = {}
    for label, o in old_pick.items():
        again = by_uid.loc[o.uid]
        same = int(again.pred) == int(o.pred) and int(again.suicide) == int(o.suicide)
        chosen[label] = again if same else new_pick[label]
        print(f"{label:20} kept original user: {same}  score {chosen[label].prob:.3f} "
              f"thr {chosen[label].threshold:.3f} pred {int(chosen[label].pred)}")

    counts = []
    for s in range(7):
        sub = new[new.suicide == s]
        counts.append({"suicide": s, "correct": int(sub.correct.sum()), "incorrect": int((1 - sub.correct).sum()),
                       "n": int(len(sub))})
    aux = pd.DataFrame(aux_fold)
    examples = [{"label": label, "suicide": int(row.suicide), "prob": float(row.prob),
                 "pred": int(row.pred), "threshold": float(row.threshold)} for label, row in chosen.items()]
    stats = {
        "model": "aux-pretrained MTM with weight averaging (scripts/boost_mtm.py, twoe), 3 seeds x top-5 epochs",
        "fold_auc": aucs,
        "counts": counts,
        "aux_mean_r": {c: float(aux[c].mean()) for c in AUX_TARGETS},
        "aux_std_r": {c: float(aux[c].std(ddof=0)) for c in AUX_TARGETS},
        "aux_fold_r": aux_fold,
        "examples": examples,
        "n_test": int(len(new)),
        "accuracy": float(new.correct.mean()),
        "thresholds": [float(t) for t in new.groupby("fold")["threshold"].first()],
    }
    (POSTER / "panel_stats_v1.json").write_text(json.dumps(stats, indent=2), encoding="utf-8")
    print("mean r without SWL:", round(float(np.mean([stats["aux_mean_r"][c] for c in AUX_TARGETS if c != "SWL"])), 3))


def write_roc() -> list[float]:
    """STM vs twoe MTM scores for the ROC panel. STM logits come from the fold cache."""
    from ssr.train.metrics import auc_roc

    runs = boost.load_runs(VARIANT)
    ys, stm, mtm, aucs = [], [], [], []
    for i in range(5):
        c = np.load(boost.CACHE / f"fold{i}.npz")
        _, te, _ = ensemble(runs, i)
        ys.append(c["yte"].astype(int))
        stm.append(sigmoid(c["stm_te"]))
        mtm.append(sigmoid(te))
        aucs.append(float(auc_roc(ys[-1], te)))
    np.savez(POSTER / "roc_scores_v1.npz",
             **{f"y{i}": ys[i] for i in range(5)},
             **{f"stm{i}": stm[i] for i in range(5)},
             **{f"mtm{i}": mtm[i] for i in range(5)})
    return aucs


def draw() -> None:
    import make_diagrams
    import make_panels
    from make_figures import STM_FOLDS
    from svgkit import MTM_C, STM_C

    stats = json.loads((POSTER / "panel_stats_v1.json").read_text(encoding="utf-8"))
    aucs = write_roc()
    print("twoe fold AUC", [round(a, 4) for a in aucs], "mean", round(float(np.mean(aucs)), 4))
    OUT.mkdir(parents=True, exist_ok=True)
    make_panels.A0 = OUT
    make_panels.draw_headline(ours=[("STM", STM_FOLDS, STM_C, 0.629), ("MTM", aucs, MTM_C, 0.697)])
    make_panels.draw_roc(True, POSTER / "roc_scores_v1.npz",
                         (("stm", STM_C, "STM"), ("mtm", MTM_C, "MTM")))
    make_panels.draw_errors(stats, True)
    make_panels.draw_aux(stats, True)
    make_panels.draw_examples_a0(stats, V1_TEXTS)
    make_diagrams.method(v1_folds=aucs, out=OUT / "method.svg")
    print("wrote", OUT)


def _set_run(run, text: str, color: str | None = None) -> None:
    run.text = text
    if color is not None:
        from pptx.dml.color import RGBColor
        run.font.color.rgb = RGBColor.from_string(color)


def patch() -> None:
    from pptx import Presentation
    from pptx.oxml.ns import qn

    pptx = POSTER / "a0" / "poster_A0_v1.pptx"
    prs = Presentation(str(pptx))
    slide = prs.slides[0]
    by_name = {sh.name: sh for sh in slide.shapes}

    tldr = by_name["TextBox 7"].text_frame.paragraphs[1].runs
    _set_run(tldr[3], "0.697 to 0.738", "0F766E")
    _set_run(tldr[4], " (multi-task).")

    _set_run(by_name["TextBox 42"].text_frame.paragraphs[0].runs[0],
             "The MTM's middle stages now track the questionnaires (mean r 0.19; dots = folds).")
    _set_run(by_name["TextBox 44"].text_frame.paragraphs[0].runs[0],
             "One held-out user per cell. Quotes are verbatim excerpts from their posts; names removed.")

    lim = by_name["TextBox 48"].text_frame.paragraphs[2].runs[0]
    _set_run(lim, "Even the true questionnaire scores only reach 0.80 AUC, so the middle stages "
             "(mean r 0.19) cannot lift the suicide head much further.")

    pictures = {
        "Picture 31": "method.svg",
        "Picture 35": "headline.svg",
        "Picture 36": "roc.svg",
        "Picture 39": "errors_by_score.svg",
        "Picture 40": "aux_scores.svg",
        "Picture 43": "examples.svg",
    }
    for name, fname in pictures.items():
        blob = (OUT / fname).read_bytes()
        el = by_name[name]._element
        rids = []
        for node in el.iter():
            rid = node.get(qn("r:embed"))
            if rid:
                rids.append(rid)
        if not rids:
            raise SystemExit(f"no image relationship on {name}")
        part = slide.part.related_part(rids[0])
        part._blob = blob
        print(f"replaced {name} <- {fname} ({len(blob)} bytes, {rids})")
    prs.save(str(pptx))
    print("wrote", pptx)


def false_alarms() -> None:
    """Print the highest-scoring false alarms of the new model with their distress-related posts (local only)."""
    import re

    cohort, folds = fold_users()
    new, _, _aucs = records(cohort, folds, "new")
    posts = pd.read_parquet(ROOT / "artifacts" / "posts.parquet", columns=["UserId", "free_text"])
    cands = new[(new.suicide == 0) & (new.pred == 1)].sort_values("prob", ascending=False).head(3)
    words = re.compile(r"\b(sad|depress\w*|alone|lonely|hate myself|tired of|cry\w*|anxi\w*|fail\w*|worthless|"
                       r"hopeless|hurt\w*|pain|give up|empty|broken|miserable|lost|numb)\b", re.I)
    for rank, (_, row) in enumerate(cands.iterrows()):
        mine = posts[posts.UserId == row.uid]["free_text"].dropna()
        hits = [t for t in mine if words.search(t)]
        print(f"=== candidate {rank}: score {row.prob:.3f}, threshold {row.threshold:.3f}, true score {row.suicide}, "
              f"{len(mine)} posts, {len(hits)} matching")
        for t in hits[:25]:
            print("  *", t[:400].replace("\n", " "))


if __name__ == "__main__":
    {"score": score, "draw": draw, "patch": patch, "false_alarms": false_alarms}[sys.argv[1]]()
